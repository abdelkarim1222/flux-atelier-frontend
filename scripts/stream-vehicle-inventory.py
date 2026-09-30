import json
import sys
from datetime import date, datetime, time
import unicodedata

from openpyxl import load_workbook


def serializable(value):
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return value


def text_value(value):
    if value is None or value == "":
        return None
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return str(value).strip()


def unique_headers(values):
    used = {}
    headers = []
    for index, value in enumerate(values, start=1):
        label = str(value).strip() if value is not None else ""
        if not label:
            label = f"Colonne {index}"
        used[label] = used.get(label, 0) + 1
        if used[label] > 1:
            label = f"{label} ({used[label]})"
        headers.append(label)
    return headers


def normalized_header(value):
    decomposed = unicodedata.normalize("NFD", value).encode("ascii", "ignore").decode("ascii")
    return " ".join(decomposed.lower().replace("°", " ").replace("º", " ").split())


def emit(record):
    print(json.dumps(record, ensure_ascii=True, separators=(",", ":")), flush=True)


def main():
    if len(sys.argv) != 2:
        raise SystemExit("Usage : python stream-vehicle-inventory.py <classeur.xlsx>")

    workbook = load_workbook(sys.argv[1], read_only=True, data_only=True)
    worksheet = workbook.active
    iterator = worksheet.iter_rows(values_only=True)
    header_row = next(iterator, None)
    if not header_row:
        raise SystemExit("Le classeur ne contient pas de ligne d’en-tête.")
    headers = unique_headers(header_row)
    normalized_headers = [normalized_header(header) for header in headers]
    if "vin" not in normalized_headers:
        raise SystemExit("La feuille Excel doit contenir une colonne « VIN ».")
    field_headers = {
        "serialNo": ["n de serie", "numero de serie", "no de serie"],
        "vin": ["vin"],
        "brandCode": ["code marque"],
        "modelCode": ["code modele"],
        "modelDescription": ["description"],
        "registration": ["n immatriculation", "numero immatriculation", "no immatriculation"],
        "stockStatus": ["stocks"],
        "warehouseCode": ["code magasin"],
        "locationCode": ["code emplacement"],
        "customerCode": ["n client", "numero client", "no client"],
        "customerName": ["nom du client"],
    }
    column_indices = {}
    for field, aliases in field_headers.items():
        column_indices[field] = next(
            (normalized_headers.index(alias) for alias in aliases if alias in normalized_headers),
            None,
        )
    emit({"type": "metadata", "sheet": worksheet.title, "headers": headers})

    imported = 0
    for excel_row, row in enumerate(iterator, start=2):
        if not any(value is not None and value != "" for value in row):
            continue
        values = list(row[:len(headers)])
        if len(values) < len(headers):
            values.extend([None] * (len(headers) - len(values)))
        raw_data = {headers[index]: serializable(value) for index, value in enumerate(values)}
        indexed_values = {
            field: text_value(values[index]) if index is not None and index < len(values) else None
            for field, index in column_indices.items()
        }
        emit({
            "type": "row",
            "sourceRow": excel_row,
            "serialNo": indexed_values["serialNo"] or f"Ligne {excel_row}",
            **indexed_values,
            "data": raw_data,
        })
        imported += 1
        if imported % 10000 == 0:
            print(f"{imported} lignes Excel lues…", file=sys.stderr, flush=True)

    workbook.close()
    print(f"Extraction terminée : {imported} lignes.", file=sys.stderr, flush=True)


if __name__ == "__main__":
    main()
